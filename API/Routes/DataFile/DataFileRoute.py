from flask import Blueprint, Response, jsonify, request, send_file, session
from pathlib import Path
import shutil, datetime, time, os,logging
from Classes.Case.DataFileClass import DataFile
from Classes.Base import Config

logger = logging.getLogger(__name__)

datafile_api = Blueprint('DataFileRoute', __name__)

@datafile_api.route("/generateDataFile", methods=['POST'])
def generateDataFile():
    try:
        casename = request.json['casename']
        caserunname = request.json['caserunname']

        if casename != None:
            txtFile = DataFile(casename)
            txtFile.generateDatafile(caserunname)
            response = {
                "message": "You have created data file!",
                "status_code": "success"
            }      
        return jsonify(response), 200
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/createCaseRun", methods=['POST'])
def createCaseRun():
    try:
        casename = request.json['casename']
        caserunname = request.json['caserunname']
        data = request.json['data']

        if casename != None:
            caserun = DataFile(casename)
            response = caserun.createCaseRun(caserunname, data)
     
        return jsonify(response), 200
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/updateCaseRun", methods=['POST'])
def updateCaseRun():
    try:
        casename = request.json['casename']
        caserunname = request.json['caserunname']
        oldcaserunname = request.json['oldcaserunname']
        data = request.json['data']

        if casename != None:
            caserun = DataFile(casename)
            response = caserun.updateCaseRun(caserunname, oldcaserunname, data)
     
        return jsonify(response), 200
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/deleteCaseRun", methods=['POST'])
def deleteCaseRun():
    try:        
        casename = request.json['casename']
        caserunname = request.json['caserunname']
        resultsOnly = request.json['resultsOnly']
        
        # casePath = Path(Config.DATA_STORAGE, casename, 'res', caserunname)
        # if not resultsOnly:
        #     shutil.rmtree(casePath)
        # else:
        #     for item in os.listdir(casePath):
        #         item_path = os.path.join(casePath, item)
        #         if os.path.isfile(item_path) or os.path.islink(item_path):
        #             os.remove(item_path)  # delete file
        #         elif os.path.isdir(item_path):
        #             shutil.rmtree(item_path)  # delete subfolder

        if casename != None:
            caserun = DataFile(casename)
            response = caserun.deleteCaseRun(caserunname, resultsOnly)    
        return jsonify(response), 200

    except(IOError):
        return jsonify('No existing cases!'), 404
    except OSError:
        raise OSError

@datafile_api.route("/deleteScenarioCaseRuns", methods=['POST'])
def deleteScenarioCaseRuns():
    try:
        scenarioId = request.json['scenarioId']
        casename = request.json['casename']

        if casename != None:
            caserun = DataFile(casename)
            response = caserun.deleteScenarioCaseRuns(scenarioId)
     
        return jsonify(response), 200
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/saveView", methods=['POST'])
def saveView():
    try:
        casename = request.json['casename']
        param = request.json['param']
        data = request.json['data']

        if casename != None:
            caserun = DataFile(casename)
            response = caserun.saveView(data, param)
     
        return jsonify(response), 200
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/updateViews", methods=['POST'])
def updateViews():
    try:
        casename = request.json['casename']
        param = request.json['param']
        data = request.json['data']

        if casename != None:
            caserun = DataFile(casename)
            response = caserun.updateViews(data, param)
     
        return jsonify(response), 200
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/readDataFile", methods=['POST'])
def readDataFile():
    try:
        casename = request.json['casename']
        caserunname = request.json['caserunname']
        if casename != None:
            txtFile = DataFile(casename)
            data = txtFile.readDataFile(caserunname)
            response = data    
        else:  
            response = None     
        return jsonify(response), 200
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/readModelFile", methods=['GET'])
def readModelFile():
    try:
        text = Path(Config.SOLVERs_FOLDER,'model.v.5.4.txt') .read_text(encoding="utf-8", errors="replace")
        return Response(text, mimetype="text/plain; charset=utf-8")
    except(IOError):
        return jsonify('No existing cases!'), 404
     
@datafile_api.route("/readLogFile", methods=['GET'])
def readLogFile():
    try:
        text = Path(Config.WebAPP_PATH,'app.log') .read_text(encoding="utf-8", errors="replace")
        return Response(text, mimetype="text/plain; charset=utf-8")
    except(IOError):
        return jsonify('No existing cases!'), 404
    
@datafile_api.route("/validateInputs", methods=['POST'])
def validateInputs():
    try:
        casename = request.json['casename']
        caserunname = request.json['caserunname']
        if casename != None:
            df = DataFile(casename)
            validation = df.validateInputs(caserunname)
            response = validation    
        else:  
            response = None     
        return jsonify(response), 200
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/downloadDataFile", methods=['GET'])
def downloadDataFile():
    try:
        #casename = request.json['casename']
        #casename = 'DEMO CASE'
        # txtFile = DataFile(casename)
        # downloadPath = txtFile.downloadDataFile()
        # response = {
        #     "message": "You have downloaded data.txt to "+ str(downloadPath) +"!",
        #     "status_code": "success"
        # }         
        # return jsonify(response), 200
        #path = "/Examples.pdf"
        case = session.get('osycase', None)
        caserunname = request.args.get('caserunname')
        dataFile = Path(Config.DATA_STORAGE,case, 'res',caserunname, 'data.txt')
        return send_file(dataFile.resolve(), as_attachment=True, max_age=0)
    
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/downloadFile", methods=['GET'])
def downloadFile():
    try:
        case = session.get('osycase', None)
        file = request.args.get('file')
        dataFile = Path(Config.DATA_STORAGE,case,'res','csv',file)
        return send_file(dataFile.resolve(), as_attachment=True, max_age=0)
    
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/downloadRunSummary", methods=['GET'])
def downloadRunSummary():
    """run_summary.csv for this case run: one row per run, with the solver, the
    HiGHS settings in force, the time of each stage and the peak memory."""
    try:
        case = session.get('osycase', None)
        caserunname = request.args.get('caserunname')
        summary = Path(Config.DATA_STORAGE, case, 'res', caserunname, 'run_summary.csv')
        if not summary.is_file():
            return jsonify('No run summary yet: run the model once.'), 404
        return send_file(summary.resolve(), as_attachment=True, max_age=0)

    except(IOError):
        return jsonify('No existing cases!'), 404


@datafile_api.route("/clearRunSummary", methods=['POST'])
def clearRunSummary():
    """Start a fresh run record for this case run.

    The existing file is renamed rather than deleted, so earlier runs are kept.
    Old records stay in the same folder as run_summary_YYYYmmdd_HHMMSS.csv and can
    be removed by hand.
    """
    try:
        case = session.get('osycase', None)
        caserunname = request.json['caserunname']
        folder = Path(Config.DATA_STORAGE, case, 'res', caserunname)
        summary = folder / 'run_summary.csv'
        if not summary.is_file():
            return jsonify({'status_code': 'success', 'archived': None,
                            'message': 'There was no record to clear.'}), 200

        stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        archive = folder / 'run_summary_{}.csv'.format(stamp)
        summary.rename(archive)
        logger.info("Run summary archived as %s", archive.name)
        return jsonify({'status_code': 'success', 'archived': archive.name,
                        'message': 'Previous runs kept as ' + archive.name}), 200

    except Exception as ex:
        return jsonify({'status_code': 'error', 'message': str(ex)}), 200


@datafile_api.route("/downloadRunSummaryXlsx", methods=['GET'])
def downloadRunSummaryXlsx():
    """The same run summary as a formatted workbook.

    Built in memory from run_summary.csv each time it is asked for, and never
    stored: a workbook kept on disk would have to be rewritten on every run, and
    would fail whenever it happened to be open in Excel. The CSV stays the record
    that is appended to; this is only a view of it.
    """
    import csv as _csv
    import io
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify('openpyxl is not installed; use the CSV instead.'), 501

    try:
        case = session.get('osycase', None)
        caserunname = request.args.get('caserunname')
        summary = Path(Config.DATA_STORAGE, case, 'res', caserunname, 'run_summary.csv')
        if not summary.is_file():
            return jsonify('No run summary yet: run the model once.'), 404

        with open(summary, newline='') as f:
            rows = list(_csv.reader(f))
        if not rows:
            return jsonify('The run summary is empty.'), 404

        wb = Workbook()
        ws = wb.active
        ws.title = 'runs'
        ws.append(rows[0])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows[1:]:
            #numbers as numbers, so the sheet can be sorted and charted directly
            ws.append([float(v) if _isNumber(v) else v for v in row])
        ws.freeze_panes = 'A2'
        for i, name in enumerate(rows[0], start=1):
            width = max([len(name)] + [len(r[i - 1]) for r in rows[1:] if len(r) >= i]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 30)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return send_file(stream, as_attachment=True, max_age=0,
                         download_name='run_summary.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    except(IOError):
        return jsonify('No existing cases!'), 404


def _isNumber(value):
    try:
        float(value)
        return value.strip() != ''
    except (TypeError, ValueError):
        return False


@datafile_api.route("/downloadCSVFile", methods=['GET'])
def downloadCSVFile():
    try:
        case = session.get('osycase', None)
        file = request.args.get('file')
        caserunname = request.args.get('caserunname')
        dataFile = Path(Config.DATA_STORAGE,case,'res',caserunname,'csv',file)
        return send_file(dataFile.resolve(), as_attachment=True, max_age=0)
    
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/downloadResultsFile", methods=['GET'])
def downloadResultsFile():
    try:
        case = session.get('osycase', None)
        caserunname = request.args.get('caserunname')
        dataFile = Path(Config.DATA_STORAGE,case, 'res', caserunname,'results.txt')
        return send_file(dataFile.resolve(), as_attachment=True, max_age=0)
    
    except(IOError):
        return jsonify('No existing cases!'), 404

@datafile_api.route("/progress", methods=['POST'])
def progress():
    """Where a run has got to. Answered on another thread while /run is still
    blocked, so the interface can show the stage and its timing as it happens."""
    from Classes.Case import RunProgressClass as Progress
    try:
        return jsonify(Progress.get(request.json['casename'],
                                    request.json['caserunname'])), 200
    except Exception:
        #progress must never break a run: an unreadable request simply reports nothing
        return jsonify({'status': 'unknown'}), 200


@datafile_api.route("/run", methods=['POST'])
def run():
    try:
        logger.info("Starting long task...")
        casename = request.json['casename']
        caserunname = request.json['caserunname']
        solver = request.json['solver']
        #optional: settings from the HiGHS panel. Absent for CBC/GLPK and for any
        #client that does not send them, in which case the solver defaults apply.
        highs_options = request.json.get('highs_options') or None
        logger.info("Starting optimization process for model -- %s -- caserun -- %s --!", casename, caserunname)
        txtFile = DataFile(casename)
        response = None
        try:
            response = txtFile.run(solver, caserunname, highs_options=highs_options)
        finally:
            #always close the progress record, including when the run raises, so the
            #interface never polls a stage that will never finish
            from Classes.Case import RunProgressClass as Progress
            Progress.finish(casename, caserunname,
                            (response or {}).get('status_code', 'error'))
            #one row per run, appended, so successive runs of the same case can be
            #compared: solver, HiGHS settings actually in force, and stage timings
            Progress.writeSummaryCsv(Path(Config.DATA_STORAGE, casename, 'res', caserunname),
                                     casename, caserunname, solver, highs_options)
        logger.info("Optimization finished for model -- %s -- caserun -- %s --!", casename, caserunname) 
        #logger.info(f"\033[92mStarting optimization process for model -- {casename} -- caserun -- {caserunname} --!\033[0m")
        return jsonify(response), 200
    # except Exception as ex:
    #     print(ex)
    #     return ex, 404
    
    except(IOError):
        return jsonify('No existing cases!'), 404
    
@datafile_api.route("/batchRun", methods=['POST'])
def batchRun():
    try:
        start = time.time()
        modelname = request.json['modelname']
        cases = request.json['cases']

        if modelname != None:
            txtFile = DataFile(modelname)
            for caserun in cases:
                logger.info("Data file generation process started for model %s caserun %s!", modelname, caserun)
                txtFile.generateDatafile(caserun)
                logger.info("Data file generation process finished for model%s caserun %s!", modelname, caserun)
            response = txtFile.batchRun( 'CBC', cases) 
        end = time.time()  
        response['time'] = end-start 
        return jsonify(response), 200
    except(IOError):
        return jsonify('Error!'), 404
    
@datafile_api.route("/cleanUp", methods=['POST'])
def cleanUp():
    try:
        modelname = request.json['modelname']

        if modelname != None:
            model = DataFile(modelname)
            logger.info("Clean up process started!")
            response = model.cleanUp()  
            logger.info("Clean up process finished!") 

        return jsonify(response), 200
    except(IOError):
        return jsonify('Error!'), 404